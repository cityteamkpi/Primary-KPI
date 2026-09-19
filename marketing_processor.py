import warnings
import os
import openpyxl
import pandas as pd
from auth_utils import get_services
from drive_utils import (
    resolve_folder_id,
    download_drive_file,
    find_file_id
)

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# ============================================================
#  CONFIGURATION & FILE NAMES
# ============================================================
TARGET_FY = "FY26"

FILE_OCCUPANCY_NAME = "Clients in All Programs Average.xlsx"
FILE_RENEW_NAME = "Renew Programs Report - Processed.xlsx"
FILE_TP_NAME = "Turning Point Report - Processed.xlsx"
FILE_MRB_NAME = "Mentorship Reunifications Baptisms - Processed.xlsx"
DESTINATION_FILE_NAME = "Marketing Impact Data.xlsx"

# Column Index Mapping in Marketing Impact Data (B=2, C=3, ..., J=10)
COLUMN_INDICES = {
    "San Jose Men": 2,      # Col B
    "San Jose Women": 3,    # Col C
    "Oakland Men": 4,       # Col D
    "Oakland Women": 5,     # Col E
    "Portland Men": 6,      # Col F
    "Portland Women": 7,    # Col G
    "Chester Men": 8,       # Col H
    "Chester Women": 9,     # Col I
    "San Francisco": 10     # Col J
}

PROGRAM_MAP = {
    "San Jose Men": {
        "renew": ["San Jose Men Renew"],
        "tp": ["San Jose Men Turning Point"]
    },
    "San Jose Women": {
        "renew": ["GV Renew"],
        "tp": ["GV Turning Point", "Heritage Home", "San Jose Youth Collective"]
    },
    "Oakland Men": {
        "renew": ["Oakland Renew"],
        "tp": ["Oakland Men Turning Point"]
    },
    "Oakland Women": {
        "renew": [],
        "tp": ["Oakland Women Turning Point"]
    },
    "Portland Men": {
        "renew": ["Portland Renew"],
        "tp": ["Portland Men Turning Point"]
    },
    "Portland Women": {
        "renew": [],
        "tp": ["Portland Community of Hope"]
    },
    "Chester Men": {
        "renew": ["Chester Renew"],
        "tp": ["Chester Men Turning Point"]
    },
    "Chester Women": {
        "renew": [],
        "tp": ["Chester Women Turning Point", "Chester Women's Turning Point"]
    },
    "San Francisco": {
        "renew": [],
        "tp": []
    }
}

ROW_RENEW_OCCUPANCY = 5
ROW_TP_OCCUPANCY = 6
ROW_RENEW_GRADUATIONS = 9
ROW_TP_GRADUATIONS = 10
ROW_TP_SUCCESSFULLY_HOUSED = 12
ROW_FAMILIES_REUNIFIED = 13
ROW_SPIRITUAL_MENTORSHIPS = 14
ROW_BAPTISMS = 15


def run_marketing_processor(
    input_folder_name="Quarterly KPI Processed Data",
    output_folder_name="Marketing Impacts"
):
    print("🚀 Starting Marketing Impact Data Processing...")

    # Authenticate seamlessly via cached token.json / auth_utils
    drive_service, gc, _ = get_services()

    # 1. Resolve Drive Folder IDs
    input_folder_id = resolve_folder_id(drive_service, input_folder_name, "Input")
    output_folder_id = resolve_folder_id(drive_service, output_folder_name, "Output")

    # 2. Download Input Excel Reports from Drive
    print("   Downloading source report files from Google Drive...")
    fh_occ, _, _ = download_drive_file(drive_service, FILE_OCCUPANCY_NAME, input_folder_id)
    fh_renew, _, _ = download_drive_file(drive_service, FILE_RENEW_NAME, input_folder_id)
    fh_tp, _, _ = download_drive_file(drive_service, FILE_TP_NAME, input_folder_id)
    fh_mrb, _, _ = download_drive_file(drive_service, FILE_MRB_NAME, input_folder_id)
    fh_dest, _, _ = download_drive_file(drive_service, DESTINATION_FILE_NAME, output_folder_id)

    # 3. Load Source DataFrames
    df_occ = pd.read_excel(fh_occ, sheet_name="Latest Occupancy")
    df_renew_grad = pd.read_excel(fh_renew, sheet_name="Renew Graduates")
    df_tp_grad = pd.read_excel(fh_tp, sheet_name="Turning Point Graduates")
    df_tp_housed = pd.read_excel(fh_tp, sheet_name="Successfully Housed")

    xls_mrb = pd.ExcelFile(fh_mrb)
    df_ment = pd.read_excel(xls_mrb, sheet_name="Mentorship Processed")
    df_reun = pd.read_excel(xls_mrb, sheet_name="Reunifications Processed")
    df_bap = pd.read_excel(xls_mrb, sheet_name="Baptisms Processed")

    # 4. Open Marketing Impact Workbook
    wb = openpyxl.load_workbook(fh_dest)
    ws = wb.active

    # 5. Populate Metrics into Form Grid
    print(f"   Calculating metrics for Target Fiscal Year: {TARGET_FY}...")
    for loc_name, col_idx in COLUMN_INDICES.items():
        renew_progs = PROGRAM_MAP[loc_name]["renew"]
        tp_progs = PROGRAM_MAP[loc_name]["tp"]
        all_progs = renew_progs + tp_progs

        # Row 5: Renew Adult Occupancy
        val_r5 = df_occ[df_occ["Program"].isin(renew_progs)]["Count"].sum()
        ws.cell(row=ROW_RENEW_OCCUPANCY, column=col_idx, value=val_r5)

        # Row 6: Turning Point Adult Occupancy
        val_r6 = df_occ[df_occ["Program"].isin(tp_progs)]["Count"].sum()
        ws.cell(row=ROW_TP_OCCUPANCY, column=col_idx, value=val_r6)

        # Row 9: Renew Graduations
        val_r9 = df_renew_grad[
            (df_renew_grad["Year"] == TARGET_FY) & 
            (df_renew_grad["Program Enrolling_2091"].isin(renew_progs))
        ]["Current Period Actuals"].sum()
        ws.cell(row=ROW_RENEW_GRADUATIONS, column=col_idx, value=val_r9)

        # Row 10: Turning Point Graduations
        val_r10 = df_tp_grad[
            (df_tp_grad["Year"] == TARGET_FY) & 
            (df_tp_grad["Program Enrolling_2091"].isin(tp_progs))
        ]["Current Period Actuals"].sum()
        ws.cell(row=ROW_TP_GRADUATIONS, column=col_idx, value=val_r10)

        # Row 12: Turning Point Successfully Housed
        val_r12 = df_tp_housed[
            (df_tp_housed["Year"] == TARGET_FY) & 
            (df_tp_housed["Program Enrolling_2091"].isin(tp_progs))
        ]["Current Period Actuals"].sum()
        ws.cell(row=ROW_TP_SUCCESSFULLY_HOUSED, column=col_idx, value=val_r12)

        # Row 13: Families Reunified
        val_r13 = df_reun[
            (df_reun["Year"] == TARGET_FY) & 
            (df_reun["Program Enrolling"].isin(all_progs))
        ]["Current Period Actuals"].sum()
        ws.cell(row=ROW_FAMILIES_REUNIFIED, column=col_idx, value=val_r13)

        # Row 14: Spiritual Mentorships
        val_r14 = df_ment[
            (df_ment["Year"] == TARGET_FY) & 
            (df_ment["Program Enrolling"].isin(all_progs))
        ]["Current Period Actuals"].sum()
        ws.cell(row=ROW_SPIRITUAL_MENTORSHIPS, column=col_idx, value=val_r14)

        # Row 15: Baptisms
        val_r15 = df_bap[
            (df_bap["Year"] == TARGET_FY) & 
            (df_bap["Program Enrolling"].isin(all_progs))
        ]["Current Period Actuals"].sum()
        ws.cell(row=ROW_BAPTISMS, column=col_idx, value=val_r15)

    # 6. Save Local Output & Upload back to Google Drive
    output_local_path = "Marketing_Impact_Data_Processed.xlsx"
    wb.save(output_local_path)

    existing_file = find_file_id(drive_service, DESTINATION_FILE_NAME, output_folder_id)
    if existing_file:
        from googleapiclient.http import MediaFileUpload
        media = MediaFileUpload(output_local_path, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        drive_service.files().update(
            fileId=existing_file['id'],
            media_body=media,
            supportsAllDrives=True
        ).execute()
        print(f"✅ Updated existing file on Google Drive: {DESTINATION_FILE_NAME}")
    else:
        from googleapiclient.http import MediaFileUpload
        file_metadata = {'name': DESTINATION_FILE_NAME, 'parents': [output_folder_id]}
        media = MediaFileUpload(output_local_path, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        drive_service.files().create(
            body=file_metadata,
            media_body=media,
            supportsAllDrives=True
        ).execute()
        print(f"✅ Created new file on Google Drive: {DESTINATION_FILE_NAME}")

    print("🎉 Marketing Processor Execution Complete!")


if __name__ == "__main__":
    run_marketing_processor()